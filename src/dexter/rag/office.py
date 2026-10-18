"""Čitanje Office fajlova bez pokretanja makroa, formula ili spoljnih veza."""

import io
from datetime import date, datetime, time
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.table import Table
from openpyxl import load_workbook

from dexter.pismo import u_latinicu

MIMES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
}
SUFFIXES = {".docx", ".xlsx"}
MAX_UNPACKED = 50 * 1024 * 1024
MAX_CELLS = 500_000


def _blocks(container, depth=0):
    if depth > 20:
        raise ValueError("Word ima previše ugnježdenih tabela.")
    for block in container.iter_inner_content():
        if isinstance(block, Table):
            for row in block.rows:
                seen, values = set(), []
                for cell in row.cells:
                    if cell._tc in seen:
                        values.append("")
                    else:
                        seen.add(cell._tc)
                        values.append(" / ".join(_blocks(cell, depth + 1)))
                if any(values):
                    yield " | ".join(values)
        elif block.text.strip():
            yield block.text.strip()


def docx_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    lines = list(_blocks(doc))
    seen = set()
    for section in doc.sections:
        for name in (
            "header",
            "first_page_header",
            "even_page_header",
            "footer",
            "first_page_footer",
            "even_page_footer",
        ):
            container = getattr(section, name)
            if container._element in seen:
                continue
            seen.add(container._element)
            parts = list(_blocks(container))
            if parts:
                label = "Zaglavlje" if "header" in name else "Podnožje"
                lines.extend([f"[{label}]", *parts])
    return "\n\n".join(lines)


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value).replace("\r", "").replace("\n", " / ")


def xlsx_text(data: bytes) -> str:
    book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    try:
        if len(book.worksheets) > 100:
            raise ValueError("Excel ima više od 100 listova; podeli fajl.")
        lines, cells = [], 0
        for sheet in book.worksheets:
            if (sheet.max_column or 0) > 1000 or (sheet.max_row or 0) > 50_000:
                raise ValueError("Excel list prelazi 50000 redova ili 1000 kolona; podeli fajl.")
            # Pogrešan dimension A1:A1 ne sme sakriti podatke kasnije u XML-u.
            sheet.reset_dimensions()
            lines.append(f"[List: {sheet.title}]")
            for i, row in enumerate(sheet.iter_rows(values_only=True), 1):
                cells += max(len(row), 1)
                if i > 50_000 or len(row) > 1000 or cells > MAX_CELLS:
                    raise ValueError(
                        "Excel prelazi ograničenje obrade; podeli fajl. Ništa nije indeksirano."
                    )
                values = [_cell(v) for v in row]
                if any(values):
                    lines.append(f"Red {i}: " + " | ".join(values))
        return "\n".join(lines)
    finally:
        book.close()


def text(name: str, data: bytes) -> str:
    with ZipFile(io.BytesIO(data)) as archive:
        infos = archive.infolist()
        if len(infos) > 10_000 or sum(info.file_size for info in infos) > MAX_UNPACKED:
            raise ValueError("Office arhiva je prevelika za obradu; podeli dokument.")
    suffix = Path(name).suffix.lower()
    if suffix == ".docx":
        result = docx_text(data)
    elif suffix == ".xlsx":
        result = xlsx_text(data)
    else:
        raise ValueError("Podržani Office formati su .docx i .xlsx.")
    return u_latinicu(result)
