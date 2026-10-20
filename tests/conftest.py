import io
import os
import re
import threading
from datetime import UTC, datetime
from email import message_from_bytes
from email.message import EmailMessage
from email.utils import parsedate_to_datetime
from types import SimpleNamespace

import pytest
from imap_tools import MailMessage

# Testovi ne smeju da zavise od pravog .env-a ni od mreže.
os.environ.setdefault("TELEGRAM_BOT_TOKEN", "test")
os.environ.setdefault("TELEGRAM_OWNER_CHAT_ID", "1")
os.environ.setdefault("ANTHROPIC_API_KEY", "test")


@pytest.fixture
def invoice_notion(monkeypatch, respx_mock):
    import json
    from uuid import uuid4

    import httpx

    from dexter.config import settings
    from dexter.services import invoices

    cfg = settings()
    base = "https://api.notion.com/v1"
    client_db, client_source, archive_db, archive_source, client_id = [
        str(uuid4()) for _ in range(5)
    ]
    values = {
        "name": "Primer izdavalac PR",
        "brand": "PRIMER",
        "address": "Primer ulica 1",
        "city": "Probni grad",
        "phone": "+381600000000",
        "pib": "000000001",
        "mb": "00000001",
        "account": "000-0000000000000-00",
        "bank": "Probna banka",
        "responsible_person": "Probna osoba",
    }
    for key, env in invoices.ISSUER_FIELDS.items():
        monkeypatch.setattr(cfg, env.lower(), values[key])
    for field, value in {
        "invoice_issuer_brand": values["brand"],
        "invoice_logo_path": "/nepostojeci-logo.png",
        "invoice_last_number": None,
        "notion_api_key": "test",
        "notion_clients_db_id": client_db,
        "notion_invoices_db_id": archive_db,
        "google_invoices_folder_id": "test-folder",
        "google_client_id": "test",
        "google_client_secret": "test",
        "google_refresh_token": "test",
    }.items():
        monkeypatch.setattr(cfg, field, value)
    for database, source in ((client_db, client_source), (archive_db, archive_source)):
        respx_mock.get(f"{base}/databases/{database}").respond(
            200, json={"data_sources": [{"id": source}]}
        )
    customer = {
        "id": client_id,
        "klijent": "Primer Studio",
        "naziv": "Primer Studio DOO",
        "adresa": "Primer adresa 2",
        "grad": "Grad klijenta",
        "pib": "000000002",
        "mb": "00000002",
    }
    client_page = {
        "id": client_id,
        "properties": {
            "Klijent": {"title": [{"plain_text": customer["klijent"]}]},
            **{
                column: {"rich_text": [{"plain_text": customer[field]}]}
                for column, field in (
                    ("Naziv za fakturu", "naziv"),
                    ("Adresa", "adresa"),
                    ("Grad", "grad"),
                    ("PIB", "pib"),
                    ("MB", "mb"),
                )
            },
        },
    }
    client_query = respx_mock.post(f"{base}/data_sources/{client_source}/query").respond(
        200, json={"results": [client_page], "has_more": False}
    )
    pages = []

    def add_invoice(number):
        entry = {
            "id": str(uuid4()),
            "url": "https://notion.so/invoice",
            "properties": {
                "Broj": {"title": [{"plain_text": number}]},
                "Datum izdavanja": {"date": {"start": "2026-10-18"}},
                "Status": {"select": {"name": "Nije plaćeno"}},
                "Iznos": {"number": 100},
            },
        }
        pages.append(entry)
        return entry

    query = respx_mock.post(f"{base}/data_sources/{archive_source}/query").mock(
        side_effect=lambda request: httpx.Response(200, json={"results": pages, "has_more": False})
    )

    def create(request):
        entry = {
            "id": str(uuid4()),
            "url": "https://notion.so/saved",
            **json.loads(request.content),
        }
        pages.append(entry)
        return httpx.Response(200, json=entry)

    create_route = respx_mock.post(f"{base}/pages").mock(side_effect=create)
    respx_mock.post("https://oauth2.googleapis.com/token").respond(
        200, json={"access_token": "test", "expires_in": 3600, "token_type": "Bearer"}
    )
    generated = respx_mock.get("https://www.googleapis.com/drive/v3/files/generateIds").mock(
        side_effect=lambda request: httpx.Response(200, json={"ids": [str(uuid4())]})
    )
    upload = respx_mock.post("https://www.googleapis.com/upload/drive/v3/files").respond(
        200,
        json={
            "id": "file-test",
            "name": "Racun.pdf",
            "webViewLink": "https://drive.google.com/file/d/file-test/view",
        },
    )
    return SimpleNamespace(
        base=base,
        client_db=client_db,
        client_source=client_source,
        archive_db=archive_db,
        source=archive_source,
        customer=customer,
        client_query=client_query,
        pages=pages,
        add_invoice=add_invoice,
        query=query,
        create=create_route,
        generated=generated,
        upload=upload,
    )


@pytest.fixture
def habits_notion(monkeypatch, respx_mock):
    import copy
    import json
    from uuid import uuid4

    import httpx

    from dexter.config import settings

    base = "https://api.notion.com/v1"
    check_db, check_source, diary_db, diary_source = [str(uuid4()) for _ in range(4)]
    cfg = settings()
    for name, value in {
        "notion_api_key": "test-notion",
        "notion_checklist_db_id": check_db,
        "notion_dnevnik_db_id": diary_db,
        "checklist_pozitivne": "Trening, Doručak",
        "checklist_izbegavanja": "Bez slatkog",
        "checklist_ciljna_stavka": "Trening",
        "checklist_cilj_nedeljno": 4,
        "timezone": "Europe/Belgrade",
    }.items():
        monkeypatch.setattr(cfg, name, value)
    for database, source in [(check_db, check_source), (diary_db, diary_source)]:
        respx_mock.get(f"{base}/databases/{database}").respond(
            200, json={"data_sources": [{"id": source}]}
        )
    pages = {}

    def page(datum="2026-10-18", diary=False, **values):
        source = diary_source if diary else check_source
        props = {
            "Dan": {"title": [{"plain_text": "Недеља"}]},
            "Datum": {"date": {"start": datum}},
        }
        if diary:
            props.update(
                {
                    "Raspoloženje": {"select": None},
                    "Energija": {"select": None},
                    "Ključna reč": {"rich_text": []},
                }
            )
        else:
            props.update(
                {
                    name: {"checkbox": values.get(name, False)}
                    for name in ("Trening", "Doručak", "Bez slatkog")
                }
            )
        result = {
            "id": str(uuid4()),
            "url": "https://notion.so/test",
            "properties": props,
            "parent": {"type": "data_source_id", "data_source_id": source},
        }
        pages[result["id"]] = result
        return result

    def query(request):
        body = json.loads(request.content)
        source = request.url.path.split("/")[-2]
        filters = body["filter"].get("and", [body["filter"]])
        result = []
        for entry in pages.values():
            if entry["parent"]["data_source_id"] != source:
                continue
            value = entry["properties"]["Datum"]["date"]["start"]
            conditions = [f["date"] for f in filters]
            if all(
                ("equals" not in f or value == f["equals"])
                and ("on_or_after" not in f or value >= f["on_or_after"])
                and ("on_or_before" not in f or value <= f["on_or_before"])
                for f in conditions
            ):
                result.append(copy.deepcopy(entry))
        return httpx.Response(200, json={"results": result, "has_more": False})

    check_query = respx_mock.post(f"{base}/data_sources/{check_source}/query").mock(
        side_effect=query
    )
    diary_query = respx_mock.post(f"{base}/data_sources/{diary_source}/query").mock(
        side_effect=query
    )

    def create(request):
        body = json.loads(request.content)
        entry = page(
            body["properties"]["Datum"]["date"]["start"],
            diary=body["parent"]["data_source_id"] == diary_source,
        )
        entry["properties"].update(body["properties"])
        return httpx.Response(200, json=entry)

    def update(request):
        entry = pages[request.url.path.split("/")[-1]]
        entry["properties"].update(json.loads(request.content)["properties"])
        return httpx.Response(200, json=entry)

    create_route = respx_mock.post(f"{base}/pages").mock(side_effect=create)
    update_route = respx_mock.patch(url__regex=base + r"/pages/[^/]+$").mock(side_effect=update)
    respx_mock.get(url__regex=base + r"/pages/[^/]+$").mock(
        side_effect=lambda request: httpx.Response(200, json=pages[request.url.path.split("/")[-1]])
    )
    return SimpleNamespace(
        base=base,
        check_db=check_db,
        check_source=check_source,
        diary_db=diary_db,
        diary_source=diary_source,
        page=page,
        pages=pages,
        check_query=check_query,
        diary_query=diary_query,
        create=create_route,
        update=update_route,
        create_handler=create,
    )


@pytest.fixture
def todo_notion(monkeypatch, respx_mock):
    from dexter.config import settings

    base = "https://api.notion.com/v1"
    database = "11111111-1111-1111-1111-111111111111"
    source = "22222222-2222-2222-2222-222222222222"
    task_id = "33333333-3333-3333-3333-333333333333"
    monkeypatch.setattr(settings(), "notion_api_key", "test-notion")
    monkeypatch.setattr(settings(), "notion_todo_db_id", database)
    monkeypatch.setattr(settings(), "weekly_task_title", "Nedeljni zadatak")
    monkeypatch.setattr(settings(), "timezone", "Europe/Belgrade")
    respx_mock.get(f"{base}/databases/{database}").respond(
        200, json={"data_sources": [{"id": source}]}
    )

    def page(title="Лични задатак", **kwargs):
        return {
            "id": task_id,
            "url": f"https://notion.so/{task_id}",
            "parent": {"type": "data_source_id", "data_source_id": source},
            "properties": {
                "Zadatak": {"title": [{"plain_text": title}]},
                "Status": {"status": {"name": "Not started"}},
                "Oblast": {"select": None},
                "Prioritet": {"select": None},
                "Rok": {"date": None},
            },
            **kwargs,
        }

    return SimpleNamespace(base=base, database=database, source=source, task_id=task_id, page=page)


@pytest.fixture
def contract_notion(monkeypatch, respx_mock, todo_notion):
    from dexter.config import settings

    base = todo_notion.base
    database = "44444444-4444-4444-4444-444444444444"
    source = "55555555-5555-5555-5555-555555555555"
    parent = "66666666-6666-6666-6666-666666666666"
    monkeypatch.setattr(settings(), "notion_tasks_db_id", database)
    metadata = {
        "object": "database",
        "id": database,
        "title": [{"plain_text": "TASK BOARD"}],
        "parent": {"type": "page_id", "page_id": parent},
        "data_sources": [{"id": source}],
    }
    respx_mock.get(f"{base}/databases/{database}").respond(200, json=metadata)
    respx_mock.get(f"{base}/pages/{parent}").respond(
        200,
        json={
            "object": "page",
            "id": parent,
            "parent": {"type": "workspace", "workspace": True},
            "properties": {"title": {"type": "title", "title": [{"plain_text": "Work"}]}},
        },
    )
    props = {"Name": {"type": "title"}, "Status": {"type": "status"}, "Rok": {"type": "date"}}
    respx_mock.get(f"{base}/data_sources/{source}").respond(200, json={"properties": props})

    def page(title="Poslovni zadatak"):
        result = todo_notion.page(title)
        result["parent"]["data_source_id"] = source
        result["properties"]["Name"] = result["properties"].pop("Zadatak")
        return result

    return SimpleNamespace(
        base=base,
        database=database,
        source=source,
        page=page,
        metadata=metadata,
        props=props,
        life=todo_notion,
    )


@pytest.fixture
async def anthropic_client(respx_mock):
    import httpx
    from anthropic import AsyncAnthropic, _base_client

    sdk_http = getattr(_base_client, "httpx2", httpx)

    def transport(request):
        response = respx_mock.handler(
            httpx.Request(request.method, str(request.url), content=request.content)
        )
        return sdk_http.Response(
            response.status_code, content=response.content, headers=response.headers
        )

    async with AsyncAnthropic(
        api_key="test",
        max_retries=0,
        http_client=sdk_http.AsyncClient(transport=sdk_http.MockTransport(transport)),
    ) as client:
        yield client


@pytest.fixture
def pdf_bytes():
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=200, height=200)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    stream = DecodedStreamObject()
    stream.set_data(
        b"BT /F1 12 Tf 20 100 Td (Dokument za proveru indeksiranja PDF fajlova i pretrage znanja.) Tj ET"
    )
    page[NameObject("/Contents")] = writer._add_object(stream)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


@pytest.fixture
def office_files():
    from docx import Document
    from openpyxl import Workbook

    word = Document()
    word.add_paragraph("Почетак документа о пројекту Орион и договореној испоруци.")
    table = word.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Рок"
    table.cell(0, 1).text = "21. октобар"
    nested = table.cell(0, 1).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Угнеждена напомена"
    word.add_paragraph("Крај документа и услови за завршетак пројекта.")
    word.sections[0].header.paragraphs[0].text = "Заглавље уговора"
    word.sections[0].footer.paragraphs[0].text = "Подножје документа"
    word_stream = io.BytesIO()
    word.save(word_stream)

    excel = Workbook()
    excel.active.title = "Приходи"
    excel.active.append(["Клијент", "Износ", "Датум", "Формула"])
    excel.active.append(["Орион", 125.5, datetime(2026, 10, 17), "=SUM(B2:B3)"])
    extra = excel.create_sheet("Резерва")
    extra.sheet_state = "hidden"
    extra.append(["Други лист са напоменом о трошковима и уговореним роковима", 0, False])
    excel_stream = io.BytesIO()
    excel.save(excel_stream)
    excel.close()
    return {"ugovor.docx": word_stream.getvalue(), "budzet.xlsx": excel_stream.getvalue()}


@pytest.fixture
def make_mail():
    def make(
        subject="Projekat",
        text="Dogovor o projektu.",
        message_id="<test@example.com>",
        when="Thu, 15 Oct 2026 10:00:00 +0200",
        html=None,
        attachments=(),
    ):
        message = EmailMessage()
        message["From"] = "Marko Petrović <marko@example.com>"
        message["To"] = "vlasnik@example.com"
        message["Subject"] = subject
        message["Date"] = when
        if message_id:
            message["Message-ID"] = message_id
        if html:
            message.set_content(html, subtype="html")
        else:
            message.set_content(text)
        for name, payload, mime in attachments:
            main, sub = mime.split("/", 1)
            message.add_attachment(payload, maintype=main, subtype=sub, filename=name)
        return message.as_bytes()

    return make


@pytest.fixture
def imap_server(monkeypatch):
    from imap_tools.imap_utf7 import utf7_decode, utf7_encode

    from dexter import config
    from dexter.services import mail
    from dexter.sources import mail as source

    for name in mail.REQUIRES:
        monkeypatch.setenv(name, "test")
    config.settings.cache_clear()
    main_thread = threading.get_ident()

    class FrozenDatetime(datetime):
        @classmethod
        def now(cls, tz=None):
            return cls(2026, 10, 15, 10, tzinfo=UTC).astimezone(tz)

    monkeypatch.setattr(source, "datetime", FrozenDatetime)

    class Server:
        state = "NONAUTH"
        selected = None
        fail_folder = None
        closed = False
        fail_login = False
        uidvalidity = 1

        def __init__(self):
            self.messages = {"INBOX": [], "Poslato": []}
            self.flags = {"INBOX": "\\Inbox", "Poslato": "\\Sent"}
            self.calls = []
            self.appended = []

        def record(self, command, *args):
            assert threading.get_ident() != main_thread
            self.calls.append((command, *args))

        def _quote(self, text):
            return f'"{text}"'

        def _simple_command(self, command, *args):
            self.record(command, *args)
            assert command in {"LOGIN", "LIST", "STATUS"}
            if command == "LOGIN" and self.fail_login:
                raise RuntimeError("IMAP prijava nije uspela")
            return "OK", []

        def _untagged_response(self, typ, data, command):
            if command == "STATUS":
                return "OK", [f'"INBOX" (UIDVALIDITY {self.uidvalidity})'.encode()]
            return "OK", [
                f'({self.flags.get(name, "")}) "/" "'.encode() + utf7_encode(name) + b'"'
                for name in self.messages
            ]

        def select(self, name, readonly=False):
            self.record("SELECT", name, readonly)
            assert readonly is True
            self.selected = utf7_decode(name).strip('"')
            if self.selected == self.fail_folder:
                raise RuntimeError("IMAP prekid foldera")
            return "OK", [str(len(self.messages[self.selected])).encode()]

        def uid(self, command, *args):
            self.record(command.upper(), *args)
            messages = self.messages[self.selected]
            if command.upper() == "SEARCH":
                criteria = args[-1].decode()
                match = re.search(r"SINCE (\d{1,2}-[A-Za-z]{3}-\d{4})", criteria)
                before_match = re.search(r"BEFORE (\d{1,2}-[A-Za-z]{3}-\d{4})", criteria)
                since = (
                    datetime.strptime(match[1], "%d-%b-%Y").date() if match else datetime.min.date()
                )
                before = (
                    datetime.strptime(before_match[1], "%d-%b-%Y").date()
                    if before_match
                    else datetime.max.date()
                )
                terms = re.findall(r'(FROM|TO|SUBJECT|BODY) "([^\"]*)"', criteria)

                def matches(raw):
                    message = MailMessage.from_bytes(raw)
                    values = {
                        "FROM": message.from_,
                        "TO": ", ".join(message.to),
                        "SUBJECT": message.subject,
                        "BODY": message.text,
                    }
                    return all(value.casefold() in values[key].casefold() for key, value in terms)

                ids = [
                    str(i)
                    for i, raw in enumerate(messages, 1)
                    if since
                    <= parsedate_to_datetime(message_from_bytes(raw)["Date"]).date()
                    < before
                    and ("UNSEEN" not in criteria or i % 2 == 0)
                    and (not re.search(r"\bSEEN\b", criteria) or i % 2 == 1)
                    and matches(raw)
                ]
                return "OK", [" ".join(ids).encode()]
            assert command.upper() == "FETCH" and "BODY.PEEK[" in args[1]
            uid = int(args[0])
            if uid > len(messages):
                return "OK", [None]
            flags = "\\Seen" if uid % 2 else ""
            raw = messages[uid - 1]
            if "BODY.PEEK[HEADER]" in args[1]:
                raw = raw.split(b"\n\n", 1)[0] + b"\n\n"
            return "OK", [(f"1 (UID {uid} FLAGS ({flags}))".encode(), raw), b")"]

        def append(self, folder, flags, when, raw):
            self.record("APPEND", folder, flags)
            name = utf7_decode(folder).strip('"')
            self.messages[name].append(raw)
            self.appended.append((name, flags, raw))
            return "OK", [b"append"]

        def logout(self):
            self.record("LOGOUT")
            self.closed = True
            return "BYE", [b"logout"]

    server = Server()
    monkeypatch.setattr(mail.MailBox, "_get_mailbox_client", lambda self: server)
    yield server
    config.settings.cache_clear()
